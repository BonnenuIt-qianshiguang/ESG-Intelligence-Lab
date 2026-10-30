"""无需原始三家报告的部署前基础测试。运行 python smoke_test.py。"""
from io import BytesIO
from pathlib import Path
import fitz
from openpyxl import load_workbook
from esg_v01_backend import process_files, excel_bytes, HEADERS

ROOT = Path(__file__).resolve().parent

def pdf(text=""):
    with fitz.open() as doc:
        page=doc.new_page()
        if text:
            page.insert_text((50,70),text)
        return doc.tobytes()

def main():
    wb=load_workbook(ROOT/"已核实样本_15条.xlsx",read_only=True,data_only=True)
    try:
        data=list(wb["已核实数据"].values)
        assert data[0] == tuple(HEADERS)
        assert len(data)==16
        assert set(row[0] for row in data[1:]) == {"奔驰","宝马","大众"}
        assert all(row[9] and row[10] and row[13] for row in data[1:])
        assert all(sum(row[0]==c for row in data[1:])==5 for c in ("奔驰","宝马","大众"))
        statuses=list(wb["文件处理状态"].values)
        assert len(statuses)==4
    finally:
        wb.close()
    print("PASS: 保存样本三家企业/15条/证据与页码/状态表")
    duplicate_content=pdf("Total energy consumption 2023: 1234 MWh")
    checks=[
        ("empty.pdf",b""),
        ("fake.txt",b"hello"),
        ("broken.pdf",b"%PDF-broken"),
        ("blank_scan.pdf",pdf()),
        ("nothing.pdf",pdf("No relevant lines here.")),
        ("candidate.pdf",duplicate_content),
        ("duplicate.pdf",duplicate_content),
    ]
    batch=process_files(checks)
    statuses={row["文件"]:row["状态"] for row in batch.files}
    assert all(statuses[x]=="处理失败" for x in ("empty.pdf","fake.txt","broken.pdf","blank_scan.pdf")),statuses
    assert statuses["nothing.pdf"]=="可读，但未找到目标指标",statuses
    assert statuses["candidate.pdf"]=="完成",statuses
    assert statuses["duplicate.pdf"]=="重复文件，已跳过",statuses
    assert not batch.records
    assert any(r["报告文件"]=="candidate.pdf" and "待人工复核" in r["核验状态"] for r in batch.candidates)
    exported=excel_bytes(batch)
    wb=load_workbook(BytesIO(exported),read_only=True,data_only=True)
    try:
        assert list(wb.sheetnames)==["说明","已核实数据","待复核候选","文件处理状态"]
        assert wb["文件处理状态"].max_row==len(checks)+1
        assert wb["待复核候选"].max_row>=2
    finally:
        wb.close()
    print("PASS: 非PDF/损坏/无文字/无关键词/候选/重复文件/Excel导出")
    return True

if __name__=="__main__":
    main()
