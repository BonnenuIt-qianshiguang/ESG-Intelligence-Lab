"""无需pytest的真实PDF与边界场景检查：python smoke_test.py。"""
from io import BytesIO
from pathlib import Path
import fitz
from openpyxl import load_workbook

from esg_v01_backend import excel_bytes, process_files


def tiny_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    output = doc.tobytes(); doc.close()
    return output


def main() -> None:
    root = Path(__file__).resolve().parent
    originals = [(p.name, p.read_bytes()) for p in [
        root / "reports/Mercedes-Benz-2023.pdf",
        root / "reports/BMW-2023.pdf",
        root / "reports/Volkswagen-2023.pdf"]]
    batch = process_files(originals)
    assert len(batch.records) == 15 and all(x["状态"] == "完成" for x in batch.files)
    expected = {("奔驰", "能源消耗总量"): (5729000, 113),
                ("宝马", "能源消耗总量"): (6380652, 309),
                ("大众", "能源消耗总量"): (20790000, 40),
                ("奔驰", "范围1 CO2排放"): (538000, 183),
                ("奔驰", "范围2 CO2排放"): (83000, 183),
                ("宝马", "范围1 CO2排放"): (713933, 307),
                ("宝马", "范围2 CO2排放"): (110141, 307),
                ("大众", "范围1 CO2排放"): (4030000, 66),
                ("大众", "范围2 CO2排放"): (1960000, 66)}
    for r in batch.records:
        if r["维度"] == "E":
            assert (r["换算值"], r["PDF页码"]) == expected[(r["企业"], r["指标"])]
            assert r["原值或内容"] in r["原文证据"]
        elif r["维度"] == "S":
            assert (r["换算值"], r["PDF页码"]) == {
                "奔驰": (166051, 187), "宝马": (154950, 314), "大众": (684025, 101)}[r["企业"]]
            assert r["原值或内容"] in r["原文证据"]
        else:
            assert r["PDF页码"] == {"奔驰":41,"宝马":142,"大众":111}[r["企业"]]
            assert "corruption" in r["原文证据"].lower()
            assert "仅表示企业披露" in r["核验状态"]
        assert r["报告文件"] and r["来源网址"] and r["提取方式"] == "已知样本专用规则"
    print("PASS 三份真实PDF、九项E、三项S、三项G的证据及页码")

    extra = [("copy.pdf", originals[0][1]), ("empty.pdf", b""),
             ("image.txt", b"hello"), ("broken.pdf", b"%PDF-this is broken"),
             ("no_metrics.pdf", tiny_pdf("Nothing related to target indicators.")),
             ("candidate.pdf", tiny_pdf("Total energy consumption 2023: 1,234 MWh")),
             ("scan.pdf", tiny_pdf(""))]
    mixed = process_files(originals + extra)
    states = {x["文件"]: x["状态"] for x in mixed.files}
    assert states["copy.pdf"] == "重复文件，已跳过"
    assert all(states[x] == "处理失败" for x in ("empty.pdf", "image.txt", "broken.pdf", "scan.pdf"))
    assert states["no_metrics.pdf"] == "可读，但未找到目标指标"
    assert states["candidate.pdf"] == "完成"
    assert len(mixed.records) == 15
    assert any(x["报告文件"] == "candidate.pdf" and x["核验状态"] == "待人工复核，不能评分" for x in mixed.candidates)
    print("PASS 多份上传、重复、空文件、非PDF、损坏PDF、扫描PDF、无指标及待核查候选")

    blob = excel_bytes(mixed)
    w = load_workbook(BytesIO(blob), read_only=True)
    assert w["已核实数据"].max_row == 16
    assert w["文件处理状态"].max_row == len(mixed.files) + 1
    assert len(excel_bytes(process_files(originals))) > 0
    a = process_files(originals)
    assert a.records == batch.records and a.files == batch.files
    print("PASS Excel内存下载内容与重复运行结果一致")


if __name__ == "__main__":
    main()
